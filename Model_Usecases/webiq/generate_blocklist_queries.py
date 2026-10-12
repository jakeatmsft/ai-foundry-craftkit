"""Generate one WebIQ query for each unique rule in the blocklist workbook."""

import csv
from collections import OrderedDict
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook


INPUT_PATH = Path("jpmc block list domains 1.xlsx")
OUTPUT_PATH = Path("blocklist_all_queries.csv")


def main() -> None:
    worksheet = load_workbook(
        INPUT_PATH, read_only=True, data_only=True
    ).active
    rules: OrderedDict[str, dict[str, object]] = OrderedDict()

    for workbook_row, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
        entry = str(values[0] or "").strip().lower().rstrip("/")
        category = str(values[2] or "").strip()
        rule = rules.setdefault(
            entry,
            {
                "blocklist_rows": [],
                "blocklist_entry": entry,
                "categories": [],
            },
        )
        rule["blocklist_rows"].append(str(workbook_row))
        if category not in rule["categories"]:
            rule["categories"].append(category)

    with OUTPUT_PATH.open("w", encoding="utf-8", newline="") as output_file:
        fieldnames = [
            "rule_id",
            "blocklist_rows",
            "blocklist_entry",
            "blocked_host",
            "blocked_path",
            "categories",
            "prompt",
        ]
        writer = csv.DictWriter(output_file, fieldnames=fieldnames)
        writer.writeheader()

        for rule_id, rule in enumerate(rules.values(), start=1):
            entry = str(rule["blocklist_entry"])
            parsed = urlsplit(f"https://{entry}")
            blocked_path = parsed.path.strip("/")
            target_kind = "webpage" if blocked_path else "website"
            target_description = "contains" if blocked_path else "is about"
            writer.writerow(
                {
                    "rule_id": rule_id,
                    "blocklist_rows": "|".join(rule["blocklist_rows"]),
                    "blocklist_entry": entry,
                    "blocked_host": parsed.hostname or "",
                    "blocked_path": blocked_path,
                    "categories": "|".join(rule["categories"]),
                    "prompt": (
                        f"Find the {target_kind} at https://{entry} "
                        f"and summarize what it {target_description}."
                    ),
                }
            )

    print(f"Wrote {len(rules)} unique blocklist queries to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
